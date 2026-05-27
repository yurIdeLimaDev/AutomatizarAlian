"""
Responsabilidade: montar o XLSX de saida do fluxo Estornos Suelane.

Diferencas em relacao ao sistema antigo:
  - O valor da coluna VALOR ESTORNO vem do XLSX original.
  - Para vendedor SUELANE, usa COM CORRETORA.
  - Para os demais vendedores, usa COMISSAO VENDEDOR.
  - Linhas com anomalias ficam em laranja e recebem motivo em OBSERVACAO.
"""

import os
import sys
from datetime import datetime
from typing import List, Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill


_PARENT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _PARENT_DIR not in sys.path:
    sys.path.insert(0, _PARENT_DIR)

from pdf_parser import PDFRecord  # noqa: E402

from estorno_suelane.matcher_suelane import (
    MatchedRecordSuelane,
    get_coluna_origem,
    get_valor_estorno,
    is_coluna_origem_encontrada,
)


_FONT_NAME = "Arial"
_COL_HEADERS = [
    "CLIENTE",
    "SITUACAO",
    "MES/ABA",
    "VEND/CORRETORA",
    "Nº APOLICE",
    "FONTE DO VALOR",
    "VALOR ESTORNO",
    "OBSERVACAO",
]
_COL_WIDTHS = [42, 12, 18, 22, 18, 22, 16, 50]

_FILL_HEADER = PatternFill("solid", start_color="D9D9D9")
_FILL_TOTAL = PatternFill("solid", start_color="FFFF00")
_FILL_WARNING = PatternFill("solid", start_color="FFCC99")

_MONTHS_PT = {
    1: "JAN", 2: "FEV", 3: "MAR", 4: "ABR",
    5: "MAI", 6: "JUN", 7: "JUL", 8: "AGO",
    9: "SET", 10: "OUT", 11: "NOV", 12: "DEZ",
}


def _safe_upper(value) -> str:
    if value is None:
        return ""
    return str(value).upper()


def _current_month_label() -> str:
    now = datetime.now()
    year_short = str(now.year)[2:]
    return f"{_MONTHS_PT[now.month]}/{year_short} ESTORNOS SUELANE"


def _build_observation(
    record: MatchedRecordSuelane,
    valor: Optional[float],
) -> tuple[List[str], bool]:
    """
    Retorna (motivos, deve_pintar_laranja).

    Esta funcao eh a regra central de anomalias do fluxo Suelane.
    """
    motivos: List[str] = []

    if not record.apolice_pdf:
        motivos.append("APOLICE AUSENTE NO PDF - MATCH FEITO APENAS PELO NOME")

    if record.match_type == "APOLICE_DIFERENTE":
        motivos.append(
            f"APOLICE DIFERENTE (PDF: {record.apolice_pdf} | "
            f"XLSX: {record.apolice_xlsx})"
        )
    elif record.match_type == "FUZZY_APOLICE_DIFERENTE":
        motivos.append(
            f"MATCH APROXIMADO E APOLICE DIFERENTE "
            f"(PDF: {record.apolice_pdf} | XLSX: {record.apolice_xlsx})"
        )
    elif record.match_type == "FUZZY":
        motivos.append("MATCH APROXIMADO PELO NOME (possivel erro de digitacao)")
    elif record.match_type == "APOLICE_AUSENTE_XLSX":
        motivos.append("APOLICE AUSENTE NA PLANILHA")
    elif record.match_type == "FUZZY_APOLICE_AUSENTE_XLSX":
        motivos.append("MATCH APROXIMADO E APOLICE AUSENTE NA PLANILHA")
    elif record.match_type != "EXATO":
        motivos.append(f"MATCH NAO EXATO ({record.match_type})")

    if record.xlsx_duplicate_count > 1:
        if record.xlsx_duplicate_value_conflict:
            motivos.append(
                "DUPLICIDADE NA PLANILHA COM VALORES DIVERGENTES "
                f"({record.xlsx_duplicate_count} linhas para o mesmo "
                "cliente/vendedor/aba/apolice)"
            )
        else:
            motivos.append(
                "DUPLICIDADE NA PLANILHA "
                f"({record.xlsx_duplicate_count} linhas para o mesmo "
                "cliente/vendedor/aba/apolice)"
            )

    fonte = get_coluna_origem(record)
    if not is_coluna_origem_encontrada(record):
        motivos.append(
            f"COLUNA '{fonte}' NAO ENCONTRADA NA ABA '{record.sheet_name}'"
        )
    elif valor is None:
        motivos.append(f"VALOR AUSENTE NA COLUNA '{fonte}'")
    elif valor == 0:
        motivos.append(f"VALOR ZERO NA COLUNA '{fonte}'")
    elif valor < 0:
        motivos.append(f"VALOR NEGATIVO NA COLUNA '{fonte}'")

    return motivos, bool(motivos)


def analyze_record_suelane(record: MatchedRecordSuelane) -> tuple[List[str], bool]:
    """
    Analise publica de anomalias usada pelo report e pela UI.

    Mantem a contagem exibida no Streamlit identica ao que sera pintado em
    laranja no XLSX.
    """
    return _build_observation(record, get_valor_estorno(record))


def _write_title_row(ws, label: str) -> None:
    cell = ws["A1"]
    cell.value = label
    cell.font = Font(name=_FONT_NAME, bold=True, size=12)
    cell.alignment = Alignment(horizontal="left")


def _write_column_headers(ws) -> None:
    for col_idx, header in enumerate(_COL_HEADERS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font = Font(name=_FONT_NAME, bold=True)
        cell.fill = _FILL_HEADER
        cell.alignment = Alignment(horizontal="center")


def _write_data_rows(ws, records: List[MatchedRecordSuelane]) -> None:
    num_cols = len(_COL_HEADERS)

    for row_idx, record in enumerate(records, start=3):
        valor = get_valor_estorno(record)
        fonte = get_coluna_origem(record)

        ws.cell(row=row_idx, column=1, value=_safe_upper(record.segurado))
        ws.cell(row=row_idx, column=2, value="ESTORNO")
        ws.cell(row=row_idx, column=3, value=_safe_upper(record.sheet_name))
        ws.cell(row=row_idx, column=4, value=_safe_upper(record.vendedor))
        ws.cell(row=row_idx, column=5, value=record.apolice_pdf)
        ws.cell(row=row_idx, column=6, value=fonte)

        valor_cell = ws.cell(row=row_idx, column=7, value=valor if valor is not None else "")
        valor_cell.number_format = 'R$ #,##0.00'

        motivos, deve_pintar = _build_observation(record, valor)
        ws.cell(row=row_idx, column=8, value=" | ".join(motivos))

        if deve_pintar:
            for col in range(1, num_cols + 1):
                ws.cell(row=row_idx, column=col).fill = _FILL_WARNING


def _write_total_row(ws, total_row: int, first_data_row: int) -> None:
    bold_font = Font(name=_FONT_NAME, bold=True)

    label_cell = ws.cell(row=total_row, column=4, value="TOTAL")
    label_cell.font = bold_font
    label_cell.fill = _FILL_TOTAL
    label_cell.alignment = Alignment(horizontal="center")

    last_data_row = total_row - 1
    total_cell = ws.cell(
        row=total_row,
        column=7,
        value=f"=SUM(G{first_data_row}:G{last_data_row})",
    )
    total_cell.font = bold_font
    total_cell.fill = _FILL_TOTAL
    total_cell.number_format = 'R$ #,##0.00'


def _set_column_widths(ws) -> None:
    col_letters = ["A", "B", "C", "D", "E", "F", "G", "H"]
    for letter, width in zip(col_letters, _COL_WIDTHS):
        ws.column_dimensions[letter].width = width


def _write_not_found_sheet(wb, not_found: List[PDFRecord]) -> None:
    ws = wb.create_sheet(title="NAO ENCONTRADOS")
    headers = ["SEGURADO", "INICIO VIGENCIA", "APOLICE"]
    widths = [50, 20, 20]

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(name=_FONT_NAME, bold=True)
        cell.fill = _FILL_HEADER
        cell.alignment = Alignment(horizontal="center")

    for row_idx, rec in enumerate(not_found, start=2):
        c1 = ws.cell(row=row_idx, column=1, value=_safe_upper(rec.segurado))
        c1.font = Font(name=_FONT_NAME)

        c2 = ws.cell(row=row_idx, column=2, value=rec.inicio_vig)
        c2.font = Font(name=_FONT_NAME)
        c2.alignment = Alignment(horizontal="center")

        c3 = ws.cell(row=row_idx, column=3, value=rec.apolice)
        c3.font = Font(name=_FONT_NAME)
        c3.alignment = Alignment(horizontal="center")

    col_letters = ["A", "B", "C"]
    for letter, width in zip(col_letters, widths):
        ws.column_dimensions[letter].width = width


def build_report_suelane(
    records: List[MatchedRecordSuelane],
    output_path: str,
    not_found: Optional[List[PDFRecord]] = None,
) -> None:
    if not records:
        raise ValueError("Nenhum registro correspondido - relatorio nao gerado.")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ESTORNOS SUELANE"

    _write_title_row(ws, _current_month_label())
    _write_column_headers(ws)

    first_data_row = 3
    _write_data_rows(ws, records)

    total_row = first_data_row + len(records)
    _write_total_row(ws, total_row, first_data_row)
    _set_column_widths(ws)

    if not_found:
        _write_not_found_sheet(wb, not_found)

    wb.save(output_path)
